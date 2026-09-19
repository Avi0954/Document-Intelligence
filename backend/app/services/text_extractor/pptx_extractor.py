from typing import List
from pptx import Presentation
from app.services.text_extractor.base import BaseTextExtractor, ExtractedChunk

class PPTXTextExtractor(BaseTextExtractor):
    def extract_chunks(self, file_path: str) -> List[ExtractedChunk]:
        chunks: List[ExtractedChunk] = []
        try:
            prs = Presentation(file_path)
            chunk_idx = 0
            for slide_idx, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                
                content = "\n".join(slide_text).strip()
                if content:
                    chunks.append(
                        ExtractedChunk(
                            content=content,
                            page_number=slide_idx + 1,  # Slide index mapped to page_number
                            chunk_index=chunk_idx
                        )
                    )
                    chunk_idx += 1

        except Exception as e:
            raise ValueError(f"Failed to extract text from PPTX file: {str(e)}")

        return chunks
